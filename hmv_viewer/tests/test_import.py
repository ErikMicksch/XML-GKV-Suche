import os
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from hmv_viewer import db, exporter, importer, queries  # noqa: E402
from hmv_viewer.model import format_nummer, nummer10  # noqa: E402
from tests.make_sample import write_sample  # noqa: E402


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.conn = db.connect(os.path.join(self.tmp, "t.db"))
        self.v1 = write_sample(os.path.join(self.tmp, "20260101_HMV.xml"), 1)
        self.v2 = write_sample(os.path.join(self.tmp, "20260201_HMV.xml"), 2)

    def tearDown(self):
        self.conn.close()

    def test_nummer(self):
        self.assertEqual(nummer10(1, 24, 1, 0, 2), "0124010002")
        self.assertEqual(format_nummer("0124010002"), "01.24.01.0002")
        self.assertEqual(format_nummer("0124010"), "01.24.01.0")
        self.assertEqual(nummer10(10, 46, 2, 1, 1), "1046021001")

    def test_first_import_is_baseline(self):
        r = importer.import_file(self.conn, self.v1)
        self.assertEqual((r.n_total, r.n_new, r.n_changed, r.n_removed), (4, 4, 0, 0))
        self.assertTrue(r.first_import)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter()), 4)

    def test_second_import_tracks_changes(self):
        importer.import_file(self.conn, self.v1)
        r = importer.import_file(self.conn, self.v2)
        self.assertEqual((r.n_new, r.n_changed, r.n_removed, r.n_unchanged, r.n_stamm_changes), (1, 1, 1, 2, 1))
        h = queries.product_history(self.conn, "0124010001")
        self.assertEqual([x["change_type"] for x in h], ["GEAENDERT", "NEU"])
        self.assertEqual(h[0]["changes"]["MERKMALE"], ["Gewicht: 4,5 kg\nLeistung: 100 VA", "Gewicht: 4,7 kg\nLeistung: 100 VA"])
        removed = queries.get_product(self.conn, "0124010002")
        self.assertEqual(removed["active"], 0)
        self.assertEqual(removed["removed_import"], 2)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter()), 4)  # aktiv: 1,3,4,5
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(status="alle")), 5)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(import_id=2, status="alle")), 3)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(import_id=2, change_type="NEU", status="alle")), 1)

    def test_reactivation_and_duplicate(self):
        importer.import_file(self.conn, self.v1)
        importer.import_file(self.conn, self.v2)
        with self.assertRaises(importer.DuplicateImport):
            importer.import_file(self.conn, self.v2)
        r = importer.import_file(self.conn, self.v1, force=True)
        self.assertEqual(r.n_reactivated, 1)
        self.assertEqual(r.n_removed, 1)  # Produkt 5 verschwindet wieder
        self.assertEqual(queries.get_product(self.conn, "0124010002")["active"], 1)

    def test_zip_source_and_date(self):
        zp = os.path.join(self.tmp, "x.zip")
        with zipfile.ZipFile(zp, "w") as z:
            z.write(self.v1, "20260101_HMV.xml")
            z.writestr("Dokumentation/readme.txt", "x")
        r = importer.import_file(self.conn, zp)
        row = self.conn.execute("SELECT * FROM imports WHERE id=?", (r.import_id,)).fetchone()
        self.assertEqual(row["source_date"], "2026-01-01")
        self.assertEqual(row["logische_version"], "2.0.0")

    def test_suspicious_import_blocked(self):
        importer.import_file(self.conn, self.v1)
        tiny = os.path.join(self.tmp, "tiny.xml")
        with open(self.v1, encoding="utf-8") as fh:
            xml = fh.read()
        import re
        xml = re.sub(r"  <hv:HMV_PRODUKT>.*?</hv:HMV_PRODUKT>\n(?=  <hv:HMV_PRODUKT>)", "", xml, count=3, flags=re.S)
        with open(tiny, "w", encoding="utf-8") as fh:
            fh.write(xml)
        with self.assertRaises(importer.SuspiciousImport):
            importer.import_file(self.conn, tiny)
        self.assertEqual(len(queries.list_versions(self.conn)), 1)  # nichts geschrieben

    def test_search_umlaut_and_number(self):
        importer.import_file(self.conn, self.v1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="ÄÖÜ")), 0)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="ä ö", in_merkmale=True)), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="01.24.01.0002")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="0124010003")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="muster")), 2)

    def test_excel_export(self):
        import xml.etree.ElementTree as ET
        import zipfile
        importer.import_file(self.conn, self.v1)
        importer.import_file(self.conn, self.v2)
        out = os.path.join(self.tmp, "e.xlsx")
        res = exporter.export_excel(self.conn, out, queries.ProductFilter(status="alle"))
        self.assertEqual(res["products"], 5)
        self.assertEqual(res["history_rows"], 3)  # V2: Merkmale geändert, Produkt 2 entfernt, Produkt 5 neu
        with zipfile.ZipFile(out) as z:
            for name in z.namelist():
                ET.fromstring(z.read(name))      # alle Teile wohlgeformt
            sheet1 = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertIn("0124010001", sheet1)
        try:
            from openpyxl import load_workbook
        except ImportError:
            return                               # optional: Gegenprüfung mit openpyxl
        wb = load_workbook(out, read_only=True)
        self.assertEqual(wb.sheetnames, ["Produkte", "Änderungshistorie", "Versionen"])
        rows = list(wb["Produkte"].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "0124010001")        # führende Null bleibt erhalten (Text)
        self.assertEqual(rows[1][1], "01.24.01.0001")
        wb.close()

    def test_single_file_bundle(self):
        import subprocess
        import shutil
        build = os.path.join(ROOT, "tools", "build_single_file.py")
        subprocess.run([sys.executable, build], check=True, capture_output=True)
        single = os.path.join(self.tmp, "app")
        os.makedirs(single)
        shutil.copy(os.path.join(ROOT, "hmv_viewer_app.py"), single)
        env = dict(os.environ, PYTHONPATH="", HMV_DB=os.path.join(single, "x.db"))
        r = subprocess.run([sys.executable, os.path.join(single, "hmv_viewer_app.py"), "import", self.v1],
                           capture_output=True, text=True, env=env, cwd=single)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("4 Produkte", r.stdout)


if __name__ == "__main__":
    unittest.main()
