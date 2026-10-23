import os
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from hmv_viewer import db, exporter, importer, queries  # noqa: E402
from hmv_viewer.model import format_nummer, nummer10  # noqa: E402
from tests.make_sample import write_sample  # noqa: E402


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.conn = db.connect(os.path.join(self.tmp, "t.db"))
        self.v1 = write_sample(os.path.join(self.tmp, "20260101_HMV.xml"), 1)
        self.v2 = write_sample(os.path.join(self.tmp, "20260201_HMV.xml"), 2)

    def tearDown(self):
        self.conn.close()

    def test_nummer(self):
        self.assertEqual(nummer10(1, 24, 1, 0, 2), "0124010002")
        self.assertEqual(format_nummer("0124010002"), "01.24.01.0002")
        self.assertEqual(format_nummer("0124010"), "01.24.01.0")
        self.assertEqual(nummer10(10, 46, 2, 1, 1), "1046021001")

    def test_first_import_is_baseline(self):
        r = importer.import_file(self.conn, self.v1)
        self.assertEqual((r.n_total, r.n_new, r.n_changed, r.n_removed), (4, 4, 0, 0))
        self.assertTrue(r.first_import)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter()), 4)

    def test_second_import_tracks_changes(self):
        importer.import_file(self.conn, self.v1)
        r = importer.import_file(self.conn, self.v2)
        self.assertEqual((r.n_new, r.n_changed, r.n_removed, r.n_unchanged, r.n_stamm_changes), (1, 1, 1, 2, 1))
        h = queries.product_history(self.conn, "0124010001")
        self.assertEqual([x["change_type"] for x in h], ["GEAENDERT", "NEU"])
        self.assertEqual(h[0]["changes"]["MERKMALE"], ["Gewicht: 4,5 kg\nLeistung: 100 VA", "Gewicht: 4,7 kg\nLeistung: 100 VA"])
        removed = queries.get_product(self.conn, "0124010002")
        self.assertEqual(removed["active"], 0)
        self.assertEqual(removed["removed_import"], 2)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter()), 4)  # aktiv: 1,3,4,5
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(status="alle")), 5)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(import_id=2, status="alle")), 3)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(import_id=2, change_type="NEU", status="alle")), 1)

    def test_reactivation_and_duplicate(self):
        importer.import_file(self.conn, self.v1)
        importer.import_file(self.conn, self.v2)
        with self.assertRaises(importer.DuplicateImport):
            importer.import_file(self.conn, self.v2)
        r = importer.import_file(self.conn, self.v1, force=True)
        self.assertEqual(r.n_reactivated, 1)
        self.assertEqual(r.n_removed, 1)  # Produkt 5 verschwindet wieder
        self.assertEqual(queries.get_product(self.conn, "0124010002")["active"], 1)

    def test_zip_source_and_date(self):
        zp = os.path.join(self.tmp, "x.zip")
        with zipfile.ZipFile(zp, "w") as z:
            z.write(self.v1, "20260101_HMV.xml")
            z.writestr("Dokumentation/readme.txt", "x")
        r = importer.import_file(self.conn, zp)
        row = self.conn.execute("SELECT * FROM imports WHERE id=?", (r.import_id,)).fetchone()
        self.assertEqual(row["source_date"], "2026-01-01")
        self.assertEqual(row["logische_version"], "2.0.0")

    def test_suspicious_import_blocked(self):
        importer.import_file(self.conn, self.v1)
        tiny = os.path.join(self.tmp, "tiny.xml")
        with open(self.v1, encoding="utf-8") as fh:
            xml = fh.read()
        import re
        xml = re.sub(r"  <hv:HMV_PRODUKT>.*?</hv:HMV_PRODUKT>\n(?=  <hv:HMV_PRODUKT>)", "", xml, count=3, flags=re.S)
        with open(tiny, "w", encoding="utf-8") as fh:
            fh.write(xml)
        with self.assertRaises(importer.SuspiciousImport):
            importer.import_file(self.conn, tiny)
        self.assertEqual(len(queries.list_versions(self.conn)), 1)  # nichts geschrieben

    def test_search_umlaut_and_number(self):
        importer.import_file(self.conn, self.v1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="ÄÖÜ")), 0)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="ä ö", in_merkmale=True)), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="01.24.01.0002")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="0124010003")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="muster")), 2)

    def test_favorites(self):
        importer.import_file(self.conn, self.v1)
        key = {"gruppe": 1, "ort": 24, "untergruppe": 1, "art": 0}
        self.assertEqual(queries.favorite_key(key), "0124010")
        self.assertEqual(queries.key_from_prefix("0124010"), key)
        from hmv_viewer.favorites import FILE_NAME, FavoriteStore
        folder = os.path.join(self.tmp, "favs")
        store = FavoriteStore(folder)
        self.assertTrue(store.toggle("0124010002"))
        self.assertTrue(store.toggle("0124"))
        store.attach(self.conn)
        # ganze Gruppe 0124 ist Favorit -> alle 4 aktiven Produkte darin zählen dazu
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(favoriten=True)), 4)
        self.assertFalse(store.toggle("0124010002"))
        store.set_note("0124010003", " Hinweis notiert ")
        store.user_name = "Meier"
        store.save()
        again = FavoriteStore(folder)
        self.assertEqual(again.user_name, "Meier")
        self.assertEqual(set(again.favs), {"0124"})
        self.assertEqual(again.notes, {"0124010003": "Hinweis notiert"})
        # kaputte Datei wird gesichert, nicht überschrieben
        with open(os.path.join(folder, FILE_NAME), "w", encoding="utf-8") as fh:
            fh.write("{kaputt")
        broken = FavoriteStore(folder)
        self.assertEqual(broken.favs, {})
        self.assertTrue(broken.recovered.exists())

    def test_shared_lists(self):
        import sqlite3
        importer.import_file(self.conn, self.v1)
        lid = queries.create_list(self.conn, "Pumpen Liste", "anna")
        with self.assertRaises(sqlite3.IntegrityError):
            queries.create_list(self.conn, "pumpen liste", "bob")      # Name eindeutig (ohne Groß/klein)
        self.assertEqual(queries.add_to_list(self.conn, lid, ["0124010001", "0124010003"], "anna"), 2)
        self.assertEqual(queries.add_to_list(self.conn, lid, ["0124010001"], "bob"), 0)   # schon drin
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(liste_id=lid)), 2)
        # zweite Verbindung (= anderer Benutzer) sieht die Liste
        other = db.connect(os.path.join(self.tmp, "t.db"))
        self.assertEqual([r["n"] for r in queries.list_lists(other)], [2])
        other.close()
        self.assertEqual([r["name"] for r in queries.lists_for_product(self.conn, "0124010001")], ["Pumpen Liste"])
        self.assertEqual(queries.remove_from_list(self.conn, lid, ["0124010003"]), 1)
        queries.rename_list(self.conn, lid, "Pumpen 2027")
        queries.delete_list(self.conn, lid)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM liste_eintrag").fetchone()[0], 0)   # Kaskade

    def test_nested_lists_and_migration(self):
        import sqlite3
        # Datenbank im Format v2.0 (Liste ohne parent_id) wird ohne Datenverlust umgestellt
        path = os.path.join(self.tmp, "alt.db")
        old = sqlite3.connect(path)
        old.executescript("CREATE TABLE liste(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL UNIQUE COLLATE NOCASE,"
                          " created_by TEXT, created_at TEXT);"
                          "CREATE TABLE liste_eintrag(liste_id INTEGER NOT NULL REFERENCES liste(id) ON DELETE CASCADE,"
                          " nummer TEXT NOT NULL, added_by TEXT, added_at TEXT, PRIMARY KEY(liste_id, nummer));"
                          "INSERT INTO liste VALUES(1, 'Alt', 'anna', '2026-10-01');"
                          "INSERT INTO liste_eintrag VALUES(1, '0124010001', 'anna', '2026-10-01');")
        old.close()
        conn = db.connect(path)
        self.assertEqual([(r["name"], r["parent_id"], r["n"]) for r in queries.list_lists(conn)], [("Alt", None, 1)])
        sub = queries.create_list(conn, "Alt", "bob", parent_id=1)          # gleicher Name auf anderer Ebene erlaubt
        queries.add_to_list(conn, sub, ["01"], "bob")                       # ganze Gruppe
        importer.import_file(conn, self.v1)
        self.assertEqual(queries.count_products(conn, queries.ProductFilter(liste_id=sub)), 4)
        self.assertEqual(queries.list_groups(conn, sub), ["01"])
        self.assertIn("Alt", [r["name"] for r in queries.lists_for_product(conn, "0124010003")])  # über Gruppe
        queries.delete_list(conn, 1)                                        # Unterliste + Einträge per Kaskade
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM liste").fetchone()[0], 0)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM liste_eintrag").fetchone()[0], 0)
        conn.close()

    def test_network_safe_journal_mode(self):
        import sqlite3
        # ältere Datenbanken (WAL) werden auf das netzlaufwerk-taugliche Rollback-Journal umgestellt
        path = os.path.join(self.tmp, "wal.db")
        old = sqlite3.connect(path)
        old.execute("PRAGMA journal_mode=WAL")
        old.close()
        conn = db.connect(path)
        self.assertEqual(conn.execute("PRAGMA journal_mode").fetchone()[0], "delete")
        self.assertFalse(db.is_read_only(path))
        conn.close()

    def test_backup_restore_and_shared_notes(self):
        importer.import_file(self.conn, self.v1)
        self.assertEqual(queries.create_user(self.conn, "  Meier "), "Meier")
        self.assertEqual(queries.create_user(self.conn, "meier"), "Meier")          # gleicher Name, keine Dublette
        with self.assertRaises(ValueError):
            queries.create_user(self.conn, "   ")
        cid = queries.add_comment(self.conn, "0124010001", "Rückruf erledigt", "Meier")
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(notiz_text="rückruf")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="rückruf")), 1)
        lid = queries.create_list(self.conn, "L", "Meier")
        queries.add_to_list(self.conn, lid, ["0124010001"], "Schulz")
        self.assertEqual(queries.list_lists(self.conn)[0]["changed_by"], "Schulz")

        backup = os.path.join(self.tmp, "sicherung.db")
        queries.backup_db(self.conn, backup)
        self.assertEqual(queries.check_backup(backup), {"versionen": 1, "produkte": 4, "listen": 1})
        queries.delete_list(self.conn, lid)
        self.assertEqual(queries.delete_comment(self.conn, cid, "Schulz"), 0)    # fremde Kommentare bleiben
        self.assertEqual(queries.delete_comment(self.conn, cid, "Meier"), 1)
        queries.restore_db(self.conn, backup)
        db.ensure_schema(self.conn)
        self.assertEqual([c["text"] for c in queries.comments_for(self.conn, "0124010001")], ["Rückruf erledigt"])
        self.assertEqual([r["name"] for r in queries.list_lists(self.conn)], ["L"])
        with open(os.path.join(self.tmp, "kein.db"), "w") as fh:
            fh.write("keine Datenbank")
        with self.assertRaises(ValueError):
            queries.check_backup(os.path.join(self.tmp, "kein.db"))

    def test_comments_log_and_number_filter(self):
        import sqlite3
        # v2.2-Datenbank mit gemeinsamer Notiz -> wird zum ersten Kommentar
        path = os.path.join(self.tmp, "v22.db")
        old = sqlite3.connect(path)
        old.executescript("CREATE TABLE notiz_gemeinsam(nummer TEXT PRIMARY KEY, text TEXT NOT NULL, changed_by TEXT,"
                          " changed_at TEXT); INSERT INTO notiz_gemeinsam VALUES('0124010001', 'alt', 'Meier', '2026-10-01T10:00:00');")
        old.close()
        conn = db.connect(path)
        importer.import_file(conn, self.v1)
        self.assertEqual([(c["text"], c["user"]) for c in queries.comments_for(conn, "0124010001")], [("alt", "Meier")])
        queries.add_comment(conn, "0124010001", "neu", "Schulz")
        self.assertEqual(queries.comment_summary(conn)["0124010001"][0], 2)
        self.assertEqual(len(queries.comments_since(conn, "2026-10-02", "Meier")), 1)
        with self.assertRaises(ValueError):
            queries.add_comment(conn, "0124010001", "  ", "Meier")
        # Verlauf einer gemeinsamen Liste
        lid = queries.create_list(conn, "L", "Meier")
        queries.add_to_list(conn, lid, ["0124010001", "0124"], "Schulz")
        queries.remove_from_list(conn, lid, ["0124"], "Meier")
        queries.rename_list(conn, lid, "L2", "Meier")
        self.assertEqual([r["aktion"] for r in queries.list_log(conn, lid)], ["umbenannt", "entfernt", "hinzugefügt", "angelegt"])
        # Nummern-Filter (zuletzt angesehen / eingefügte Nummernliste)
        f = queries.ProductFilter(status="alle", nummern=("0124010003", "9999999999"))
        self.assertEqual([r["nummer"] for r in queries.fetch_products(conn, f)], ["0124010003"])
        conn.close()

    def test_regrouping_and_filters(self):
        import json
        import sqlite3
        importer.import_file(self.conn, self.v1)
        v3 = write_sample(os.path.join(self.tmp, "20260301_HMV.xml"), 3)
        r = importer.import_file(self.conn, v3)
        self.assertEqual((r.n_regrouped, r.n_new, r.n_removed), (1, 1, 1))     # Platzhalter 4 -> 8 bleibt neu/entfernt
        changes = self.conn.execute("SELECT changes FROM history WHERE k='0124010007'").fetchone()[0]
        self.assertEqual(json.loads(changes), {"NUMMER": ["01.24.01.0002", "01.24.01.0007"]})
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(status="alle", change_type="UMGRUPPIERT")), 2)
        self.assertEqual(queries.count_changes(self.conn, r.import_id, "UMGRUPPIERT"), 2)
        # „Nicht besetzt“ ausblenden
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(ohne_platzhalter=True)), 3)
        # einmalige Nachbearbeitung älterer Datenbanken (Spalte fehlt, History noch NEU/ENTFERNT)
        self.conn.execute("UPDATE history SET change_type='NEU', changes=NULL WHERE k='0124010007'")
        self.conn.execute("UPDATE history SET change_type='ENTFERNT', changes=NULL WHERE k='0124010002' AND import_id=2")
        self.conn.execute("ALTER TABLE imports DROP COLUMN n_regrouped")
        self.conn.close()
        self.conn = db.connect(os.path.join(self.tmp, "t.db"))
        self.assertEqual(queries.list_versions(self.conn)[0]["n_regrouped"], 1)
        self.assertEqual(queries.count_changes(self.conn, 2, "UMGRUPPIERT"), 2)

    def test_changed_merkmale_only(self):
        importer.import_file(self.conn, self.v1)
        importer.import_file(self.conn, self.v2)   # Pumpe A: nur Merkmale geändert
        # Pumpe C: nur Bezeichnung geändert -> darf bei „Geändert – Merkmale“ NICHT erscheinen
        self.conn.execute("INSERT INTO history(import_id, entity, k, change_type, changes) VALUES"
                          " (2, 'PRODUKT', '0124010003', 'GEAENDERT', '{\"BEZEICHNUNG\": [\"a\", \"b\"]}')")
        self.assertEqual(queries.count_changes(self.conn, 2, "GEAENDERT"), 3)   # Gruppe, Pumpe A, Pumpe C
        f = queries.ProductFilter(status="alle", change_type=queries.MERKMALE_GEAENDERT)
        self.assertEqual([p["nummer"] for p in queries.fetch_products(self.conn, f)], ["0124010001"])
        self.assertEqual(queries.count_changes(self.conn, 2, queries.MERKMALE_GEAENDERT), 1)

    def test_fulltext_and_properties(self):
        from hmv_viewer import merkmale
        importer.import_file(self.conn, self.v1)
        f = lambda text, **kw: queries.count_products(self.conn, queries.ProductFilter(**{"text": text, "in_merkmale": True, **kw}))
        self.assertEqual(f("gewicht 4,5 KG"), 1)              # Satzzeichen/Groß-klein egal
        self.assertEqual(f("4.5kg"), 1)                       # „4.5kg“ = „4,5 kg“
        self.assertEqual(f("gewi"), 1)                        # Wortanfang genügt
        self.assertEqual(f("10"), 0)                          # Zahl nur als ganzes Wort (kein Treffer in „100 VA“)
        self.assertEqual(f("100"), 1)
        self.assertEqual(f('"Leistung 100"'), 1)              # Phrase
        self.assertEqual(f('"100 Leistung"'), 0)
        self.assertEqual(f("pumpe -gewicht"), 2)              # Pumpe B und C, nicht A
        self.assertEqual(f("01.24.01 pumpe"), 3)              # Nummernbereich + Wort (Pumpe A, B, C)
        self.assertEqual(f("gewicht", in_merkmale=False), 0)  # Suchbereich Merkmale aus
        self.assertEqual(f('"Leistung: 100 VA"'), 1)          # so sucht ein Klick auf eine Merkmal-Zeile
        # Index wird beim Import aktualisiert
        importer.import_file(self.conn, self.v2)
        self.assertEqual(f("4,7 kg"), 1)
        self.assertEqual(f("4,5 kg"), 0)
        # Silbentrennung und Phrase-Erkennung
        self.assertEqual(merkmale.normalize("hö-\nhenverstellbar PR70"), "höhenverstellbar PR 70")
        self.assertEqual(merkmale.parse_query('"Typ 10" -netz 01.24').numbers, ["01.24"])
        # Fundstelle trotz doppelter Leerzeichen im Originaltext (echter Fall aus der HMV)
        self.assertEqual(merkmale.fundstelle("Bauform: HdO\nEnergiequelle:  Typ 312 (PR41)", ["energiequelle: typ 312 (PR 41)"]),
                         "Energiequelle: Typ 312 (PR41)")

    def test_others_activity(self):
        importer.import_file(self.conn, self.v1)
        queries.add_comment(self.conn, "0124010001", "von Schulz", "Schulz")
        queries.add_comment(self.conn, "0124010003", "von mir", "Meier")
        queries.log_activity(self.conn, "Schulz", "kategorie", ["0124010004"], "Kritisch", "#D64541")
        lid = queries.create_list(self.conn, "Liste Schulz", "Schulz")
        rows = queries.others_activity(self.conn, "Meier")
        self.assertEqual({r["art"] for r in rows}, {"kommentar", "kategorie", "liste"})
        self.assertNotIn("von mir", [r["detail"] for r in rows])                     # eigene Einträge nicht
        self.assertEqual([r["bezeichnung"] is not None for r in rows if r["art"] == "kategorie"], [True])
        self.assertEqual([r["liste_id"] for r in rows if r["art"] == "liste"], [lid])
        self.assertEqual(len(queries.others_activity(self.conn, "Meier", kind="kommentar")), 1)
        # jede Art einzeln liefert dieselben Spaltennamen (erster Teil der UNION bestimmt sie)
        self.assertEqual(queries.others_activity(self.conn, "Meier", kind="kategorie")[0]["detail"], "Kritisch")
        self.assertEqual(queries.others_activity(self.conn, "Meier", kind="liste")[0]["liste_id"], lid)
        self.assertEqual(queries.others_activity(self.conn, "Meier", user="Niemand"), [])
        self.assertEqual(queries.others_activity(self.conn, "Meier", since="2999-01-01"), [])

    def test_private_nested_lists(self):
        from hmv_viewer.favorites import FavoriteStore
        importer.import_file(self.conn, self.v1)
        store = FavoriteStore(os.path.join(self.tmp, "favs"))
        top = store.create_plist("Projekte")
        sub = store.create_plist("Pumpen", parent=top)
        with self.assertRaises(ValueError):
            store.create_plist("pumpen", parent=top)                       # Name je Ebene eindeutig
        store.add_to_plist(sub, ["0124", "0124010001"])
        store.attach(self.conn)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(pliste_id=sub)), 4)
        again = FavoriteStore(os.path.join(self.tmp, "favs"))
        self.assertEqual(again.plists[sub]["parent"], top)
        again.delete_plist(top)                                             # samt Unterliste
        self.assertEqual(again.plists, {})

    def test_notes_search_and_categories(self):
        from hmv_viewer.favorites import FavoriteStore
        importer.import_file(self.conn, self.v1)
        store = FavoriteStore(os.path.join(self.tmp, "favs"))
        self.assertEqual(len(store.categories), 3)                      # Vorgaben für neue Benutzer
        store.set_note("0124010003", "Lager Süd")
        cid = store.add_category("Projekt X", "#123456")
        store.set_mark(["0124010001", "0124010004"], cid)
        store.attach(self.conn)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(text="lager")), 1)
        # Feld „Merkmale/Notizen“: eine Zeile aus den Details findet das Produkt (auch ohne Merkmal-Häkchen)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(notiz_text="gewicht: 4,5 KG")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(notiz_text="lager süd")), 1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(kategorie=cid)), 2)
        again = FavoriteStore(os.path.join(self.tmp, "favs"))
        self.assertEqual(again.categories[cid], {"name": "Projekt X", "farbe": "#123456"})
        again.delete_category(cid)                                         # Markierungen verschwinden mit
        self.assertEqual(again.marks, {})

    def test_hersteller_filter(self):
        importer.import_file(self.conn, self.v1)
        self.assertEqual(queries.count_products(self.conn, queries.ProductFilter(hersteller="muster gm")), 2)
        self.assertIn("Muster GmbH", queries.list_hersteller(self.conn))

    def test_excel_export(self):
        import xml.etree.ElementTree as ET
        import zipfile
        importer.import_file(self.conn, self.v1)
        importer.import_file(self.conn, self.v2)
        out = os.path.join(self.tmp, "e.xlsx")
        res = exporter.export_excel(self.conn, out, queries.ProductFilter(status="alle"))
        self.assertEqual(res["products"], 5)
        self.assertEqual(res["history_rows"], 3)  # V2: Merkmale geändert, Produkt 2 entfernt, Produkt 5 neu
        with zipfile.ZipFile(out) as z:
            for name in z.namelist():
                ET.fromstring(z.read(name))      # alle Teile wohlgeformt
            sheet1 = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertIn("0124010001", sheet1)
        try:
            from openpyxl import load_workbook
        except ImportError:
            return                               # optional: Gegenprüfung mit openpyxl
        wb = load_workbook(out, read_only=True)
        self.assertEqual(wb.sheetnames, ["Produkte", "Änderungshistorie", "Versionen"])
        rows = list(wb["Produkte"].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "0124010001")        # führende Null bleibt erhalten (Text)
        self.assertEqual(rows[1][1], "01.24.01.0001")
        wb.close()

    def test_single_file_bundle(self):
        import subprocess
        import shutil
        build = os.path.join(ROOT, "tools", "build_single_file.py")
        if not os.path.exists(build):
            self.skipTest("Testordner ohne Build-Skript (Einzeldatei gibt es nur im Hauptprogramm)")
        subprocess.run([sys.executable, build], check=True, capture_output=True)
        single = os.path.join(self.tmp, "app")
        os.makedirs(single)
        shutil.copy(os.path.join(ROOT, "hmv_viewer_app.py"), single)
        env = dict(os.environ, PYTHONPATH="", HMV_DB=os.path.join(single, "x.db"))
        r = subprocess.run([sys.executable, os.path.join(single, "hmv_viewer_app.py"), "import", self.v1],
                           capture_output=True, text=True, env=env, cwd=single)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("4 Produkte", r.stdout)


if __name__ == "__main__":
    unittest.main()
